#%%
import docx
import pandas as pd
df =pd.read_csv('data/dc_data.csv', skiprows=2, nrows=96)

df=df.set_index('Time Block')
#df= df[['Time Block','KAPS(0)']]
#%%
# set time block as index
#df=df.set_index('Time Block')
#df= df[['Time Block','KAPS(0)']]

# open an existing document
doc= docx.Document('doc_tempate.docx')

# add a tabe at the end and create a reference variable
# add extra row for header
t= doc.add_table(df.shape[0]+1, df.shape[1])

#%%
# add the header rows
for j in range(df.shape[-1]):
    t.cell(0,j).text = df.columns[j]

#add the rest of data frame
for i in range(df.shape[0]):
    for j in range(df.shape[-1]):
        t.cell(i+1,j).text = str(df.values[i,j])

# save
doc.save('doc_tempate.docx')
# %%


# %%
